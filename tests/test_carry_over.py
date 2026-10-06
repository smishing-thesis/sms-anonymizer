import csv
import glob

import pytest
from openpyxl import Workbook, load_workbook

from sms_anonymizer.anonymize import anonymizer
from sms_anonymizer.cli import build_parser
from sms_anonymizer.export.excel_writer import write_labeling_workbook
from sms_anonymizer.label_import.carry_over import (
    LabelConflictError,
    backup_before_overwrite,
    has_labels,
    load_labels_by_text,
)


def _labeled_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "text", "label"])
    for row in rows:
        ws.append(list(row))
    wb.save(str(path))


def test_load_labels_by_text_skips_unlabeled_and_invalid_rows(tmp_path):
    path = tmp_path / "prev.xlsx"
    _labeled_workbook(
        path, [("a:0", "uno", "ham"), ("a:1", "dos", None), ("a:2", "tres", "otro"), ("a:3", "cuatro", "scam")]
    )
    assert load_labels_by_text([str(path)]) == {"uno": "ham", "cuatro": "scam"}


def test_load_labels_by_text_merges_several_workbooks(tmp_path):
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _labeled_workbook(first, [("a:0", "uno", "ham")])
    _labeled_workbook(second, [("b:0", "dos", "scam"), ("b:1", "uno", "ham")])
    assert load_labels_by_text([str(first), str(second)]) == {"uno": "ham", "dos": "scam"}


def test_conflicting_labels_for_the_same_text_raise(tmp_path):
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _labeled_workbook(first, [("a:0", "uno", "ham")])
    _labeled_workbook(second, [("b:0", "uno", "scam")])
    with pytest.raises(LabelConflictError):
        load_labels_by_text([str(first), str(second)])


def test_matching_is_exact_on_text(tmp_path):
    path = tmp_path / "out.xlsx"
    write_labeling_workbook(
        [("n:0", "Hola"), ("n:1", "hola")], str(path), existing_labels={"Hola": "scam"}
    )
    ws = load_workbook(str(path)).active
    assert [r[2].value for r in ws.iter_rows(min_row=2)] == ["scam", None]


def test_has_labels_and_backup(tmp_path):
    empty, labeled = tmp_path / "empty.xlsx", tmp_path / "labeled.xlsx"
    write_labeling_workbook([("a:0", "uno")], str(empty))
    _labeled_workbook(labeled, [("a:0", "uno", "ham")])
    assert not has_labels(str(empty))
    assert has_labels(str(labeled))

    backup = backup_before_overwrite(str(labeled))
    assert backup != str(labeled)
    assert has_labels(backup)


def _write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["text", "sender"])
        writer.writerows(rows)


def _process(tmp_path, csv_files, output, extra=()):
    argv = ["process"]
    for csv_file in csv_files:
        argv += ["--source", f"csv:{csv_file}"]
    argv += [
        "--output", str(output),
        "--metadata", str(tmp_path / "metadata.json"),
        "--stats", str(tmp_path / "stats.json"),
        "--salt-path", str(tmp_path / ".salt"),
        "--known-names", str(tmp_path / "no_names.txt"),
        *extra,
    ]
    args = build_parser().parse_args(argv)
    args.func(args)


def _labels(path):
    ws = load_workbook(str(path)).active
    return {r[1].value: r[2].value for r in ws.iter_rows(min_row=2)}


def test_new_iteration_over_the_same_output_keeps_existing_labels(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(anonymizer, "find_names", lambda text: [])
    (tmp_path / "no_names.txt").write_text("", encoding="utf-8")
    first_csv, second_csv = tmp_path / "first.csv", tmp_path / "second.csv"
    _write_csv(first_csv, [["Mensaje legitimo uno", "80800"], ["Mensaje sospechoso dos", "80800"]])
    _write_csv(second_csv, [["Mensaje totalmente nuevo", "80800"]])
    output = tmp_path / "to_label.xlsx"

    _process(tmp_path, [first_csv], output)
    ws = load_workbook(str(output)).active
    ws["C2"], ws["C3"] = "ham", "scam"  # manual labeling
    ws.parent.save(str(output))

    _process(tmp_path, [first_csv, second_csv], output)  # second iteration, more data

    assert _labels(output) == {
        "Mensaje legitimo uno": "ham",
        "Mensaje sospechoso dos": "scam",
        "Mensaje totalmente nuevo": None,
    }
    assert len(glob.glob(str(tmp_path / "to_label.bak-*.xlsx"))) == 1
    out = capsys.readouterr().out
    assert "2 keep their existing label, 1 still need labeling" in out


def test_carry_labels_from_a_different_workbook(tmp_path, monkeypatch):
    monkeypatch.setattr(anonymizer, "find_names", lambda text: [])
    (tmp_path / "no_names.txt").write_text("", encoding="utf-8")
    csv_path = tmp_path / "data.csv"
    _write_csv(csv_path, [["Mensaje legitimo uno", "80800"], ["Mensaje nuevo dos", "80800"]])
    previous = tmp_path / "previous.xlsx"
    _labeled_workbook(previous, [("old:0", "Mensaje legitimo uno", "ham")])
    output = tmp_path / "to_label_v3.xlsx"

    _process(tmp_path, [csv_path], output, extra=["--carry-labels", str(previous)])

    assert _labels(output) == {"Mensaje legitimo uno": "ham", "Mensaje nuevo dos": None}
    assert not glob.glob(str(tmp_path / "to_label_v3.bak-*.xlsx"))  # nothing to back up


def test_warns_about_previously_labeled_rows_missing_from_the_new_run(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(anonymizer, "find_names", lambda text: [])
    (tmp_path / "no_names.txt").write_text("", encoding="utf-8")
    csv_path = tmp_path / "data.csv"
    _write_csv(csv_path, [["Mensaje que sigue", "80800"]])
    previous = tmp_path / "previous.xlsx"
    _labeled_workbook(previous, [("o:0", "Mensaje que sigue", "ham"), ("o:1", "Mensaje que ya no esta", "scam")])

    _process(tmp_path, [csv_path], tmp_path / "out.xlsx", extra=["--carry-labels", str(previous)])

    assert "WARNING: 1 previously labeled" in capsys.readouterr().out
