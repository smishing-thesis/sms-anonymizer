from openpyxl import load_workbook

from sms_anonymizer.export.excel_writer import write_labeling_workbook


def test_writes_header_and_rows(tmp_path):
    output = tmp_path / "to_label.xlsx"
    write_labeling_workbook([("android_xml:0", "Hola <NAMED_ENTITY>")], str(output))

    wb = load_workbook(str(output))
    ws = wb.active

    assert [c.value for c in ws[1]] == ["id", "text", "label"]
    assert [c.value for c in ws[2]] == ["android_xml:0", "Hola <NAMED_ENTITY>", None]


def test_label_column_has_dropdown_validation(tmp_path):
    output = tmp_path / "to_label.xlsx"
    write_labeling_workbook([("id:0", "text")], str(output))

    wb = load_workbook(str(output))
    ws = wb.active

    assert len(ws.data_validations.dataValidation) == 1
    dv = ws.data_validations.dataValidation[0]
    assert dv.formula1 == '"scam,ham"'
