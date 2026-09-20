import csv
import json

from openpyxl import Workbook, load_workbook

from sms_anonymizer.anonymize import anonymizer
from sms_anonymizer.cli import build_parser


def _write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["text", "sender"])
        writer.writerows(rows)


def test_pipeline_then_finalize_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setattr(anonymizer, "find_names", lambda text: [])

    csv_path = tmp_path / "aportes.csv"
    _write_csv(csv_path, [["Escribime al 987654321", "80800"], ["Otro mensaje distinto", "80800"]])

    to_label = tmp_path / "to_label.xlsx"
    metadata_path = tmp_path / "metadata.json"
    stats_path = tmp_path / "stats.json"
    spec_path = tmp_path / "spec.json"
    verify_report = tmp_path / "verify_report.json"

    parser = build_parser()
    args = parser.parse_args(
        [
            "pipeline",
            "--spec-output",
            str(spec_path),
            "--source",
            f"csv:{csv_path}",
            "--output",
            str(to_label),
            "--metadata",
            str(metadata_path),
            "--stats",
            str(stats_path),
            "--salt-path",
            str(tmp_path / ".sender_salt"),
            "--verify-report",
            str(verify_report),
        ]
    )
    from sms_anonymizer.cli import _cmd_pipeline

    _cmd_pipeline(args)

    assert spec_path.exists()
    assert to_label.exists()
    with open(metadata_path, encoding="utf-8") as f:
        metadata = json.load(f)
    assert len(metadata) == 2
    assert metadata[0]["sender_type"] == "short_code"

    with open(verify_report, encoding="utf-8") as f:
        report = json.load(f)
    assert report["findings"] == []

    # simulate manual labeling
    wb = load_workbook(str(to_label))
    ws = wb.active
    ws["C2"] = "scam"
    ws["C3"] = "ham"
    labeled_path = tmp_path / "labeled.xlsx"
    wb.save(str(labeled_path))

    final_csv = tmp_path / "final.csv"
    manifest_path = tmp_path / "manifest.json"
    version_state = tmp_path / "version_state.json"

    finalize_args = parser.parse_args(
        [
            "finalize",
            "--input",
            str(labeled_path),
            "--output",
            str(final_csv),
            "--manifest",
            str(manifest_path),
            "--version-state",
            str(version_state),
            "--stats",
            str(stats_path),
        ]
    )
    finalize_args.func(finalize_args)

    with open(final_csv, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == ["text", "label"]
    assert rows[1][1] == "1"
    assert rows[2][1] == "0"

    with open(manifest_path, encoding="utf-8") as f:
        manifest = json.load(f)
    assert manifest["corpus_version"] == "1.0.0"
    assert manifest["replacement_counts"]["<PHONE_NUMBER>"] == 1
