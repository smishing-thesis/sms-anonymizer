from openpyxl import load_workbook

from ..schema.labels import LABEL_TO_INT


class LabelValidationError(ValueError):
    pass


def read_labeled_rows(input_path: str) -> list[tuple[str, str, str]]:
    """Returns (id, text, label) tuples. Raises if any row with text has a
    missing or invalid label, naming the offending Excel row numbers."""
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = wb.active

    result = []
    errors = []
    for excel_row_number, (row_id, text, label) in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if text is None:
            continue
        if label not in LABEL_TO_INT:
            errors.append(f"row {excel_row_number}: invalid label {label!r}")
            continue
        result.append((row_id, text, label))

    if errors:
        raise LabelValidationError("; ".join(errors))
    return result
