import os
import shutil
from datetime import datetime

from openpyxl import load_workbook

from ..schema.labels import LABEL_TO_INT


class LabelConflictError(ValueError):
    pass


def load_labels_by_text(paths: list[str]) -> dict[str, str]:
    """Read already-assigned labels from labeling workbooks, keyed by the exact
    anonymized text. Rows without a (valid) label are skipped. The same text
    labeled differently in two places raises, rather than silently picking one."""
    labels: dict[str, str] = {}
    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for excel_row_number, (_row_id, text, label) in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            if text is None or label not in LABEL_TO_INT:
                continue
            if labels.setdefault(text, label) != label:
                raise LabelConflictError(
                    f"{path} row {excel_row_number}: same text already labeled "
                    f"{labels[text]!r}, now {label!r}: {text[:80]!r}"
                )
        wb.close()
    return labels


def has_labels(path: str) -> bool:
    return bool(load_labels_by_text([path]))


def backup_before_overwrite(path: str) -> str:
    """Copy `path` next to itself with a timestamp, e.g. messages.bak-20261006-152300.xlsx."""
    root, ext = os.path.splitext(path)
    backup_path = f"{root}.bak-{datetime.now():%Y%m%d-%H%M%S}{ext}"
    shutil.copy2(path, backup_path)
    return backup_path
