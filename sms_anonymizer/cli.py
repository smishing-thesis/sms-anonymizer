import argparse
import json

from openpyxl import load_workbook

from .anonymize.spec_export import write_spec
from .export.excel_writer import write_labeling_workbook
from .ingest.registry import PARSERS
from .label_import.csv_writer import write_final_csv
from .label_import.excel_reader import read_labeled_rows
from .manifest.manifest_writer import write_manifest
from .pipeline.orchestrator import run_ingest_and_anonymize
from .verify.residual_scan import scan_residuals
from .verify.review_sample import write_review_sample


def _parse_source(raw: str) -> tuple[str, str]:
    if ":" not in raw:
        raise argparse.ArgumentTypeError(f"--source must be FORMAT:PATH, got {raw!r}")
    format_name, path = raw.split(":", 1)
    if format_name not in PARSERS:
        raise argparse.ArgumentTypeError(f"Unknown format {format_name!r}. Known: {sorted(PARSERS)}")
    return format_name, path


def _read_id_text_pairs(path: str) -> list[tuple[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return [(row[0], row[1]) for row in ws.iter_rows(min_row=2, values_only=True) if row[1] is not None]


def _cmd_spec_export(args: argparse.Namespace) -> None:
    write_spec(args.output)
    print(f"Wrote anonymization rule spec to {args.output}")


def _cmd_process(args: argparse.Namespace) -> None:
    sources = [_parse_source(s) for s in args.source]
    messages, before_after, counts_by_source, discards, replacements = run_ingest_and_anonymize(
        sources, salt_path=args.salt_path, min_length=args.min_length
    )

    write_labeling_workbook([(m.id, m.text) for m in messages], args.output)
    print(f"Wrote {len(messages)} anonymized messages to {args.output}")

    metadata = [
        {
            "id": m.id,
            "source": m.source,
            "service": m.service,
            "sender_type": m.sender_type,
            "sender_pseudo_id": m.sender_pseudo_id,
            "message_timestamp": m.message_timestamp,
            "processed_at": m.processed_at,
        }
        for m in messages
    ]
    with open(args.metadata, "w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False, indent=2)

    stats = {
        "counts_by_source": counts_by_source,
        "discards_by_reason": discards,
        "replacement_counts": replacements,
    }
    with open(args.stats, "w", encoding="utf-8") as f:
        json.dump(stats, f, ensure_ascii=False, indent=2)

    if args.review_sample:
        sample_count = write_review_sample(
            before_after, args.review_sample, args.review_sample_size, seed=args.review_sample_seed
        )
        print(f"Wrote a review sample of {sample_count} messages to {args.review_sample}")


def _cmd_verify(args: argparse.Namespace) -> None:
    pairs = _read_id_text_pairs(args.input)
    findings = scan_residuals(pairs)
    with open(args.report, "w", encoding="utf-8") as f:
        json.dump({"scanned": len(pairs), "findings": findings}, f, ensure_ascii=False, indent=2)
    print(f"Scanned {len(pairs)} messages, {len(findings)} with residual PII. Report: {args.report}")


def _cmd_finalize(args: argparse.Namespace) -> None:
    rows = read_labeled_rows(args.input)
    write_final_csv(rows, args.output)
    print(f"Wrote {len(rows)} labeled rows to {args.output}")

    stats = {"counts_by_source": {}, "discards_by_reason": {}, "replacement_counts": {}}
    if args.stats:
        with open(args.stats, encoding="utf-8") as f:
            stats = json.load(f)

    manifest = write_manifest(
        output_file_path=args.output,
        manifest_path=args.manifest,
        version_state_path=args.version_state,
        counts_by_source=stats["counts_by_source"],
        discards_by_reason=stats["discards_by_reason"],
        replacement_counts=stats["replacement_counts"],
    )
    print(f"Corpus version {manifest['corpus_version']} — manifest written to {args.manifest}")


def _add_process_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument(
        "--source",
        action="append",
        required=True,
        help=f"FORMAT:PATH, repeatable. Known formats: {sorted(PARSERS)}",
    )
    parser.add_argument("--output", required=True, help="Excel file for manual scam/ham labeling")
    parser.add_argument("--metadata", required=True, help="JSON sidecar with per-message traceability")
    parser.add_argument("--stats", required=True, help="JSON sidecar with run counts, for the manifest")
    parser.add_argument("--salt-path", default="data/.sender_salt")
    parser.add_argument("--min-length", type=int, default=3)
    parser.add_argument("--review-sample", help="Optional CSV path for a random before/after review sample")
    parser.add_argument("--review-sample-size", type=int, default=30)
    parser.add_argument("--review-sample-seed", type=int, default=None)
    parser.set_defaults(func=_cmd_process)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="sms-anonymizer")
    sub = parser.add_subparsers(dest="command", required=True)

    p_spec = sub.add_parser("spec-export", help="Write the anonymization rule spec as JSON")
    p_spec.add_argument("--output", required=True)
    p_spec.set_defaults(func=_cmd_spec_export)

    p_process = sub.add_parser(
        "process", help="Ingest raw messages, clean, anonymize, dedupe, export for labeling"
    )
    _add_process_args(p_process)

    p_verify = sub.add_parser("verify", help="Scan already-anonymized output for residual PII")
    p_verify.add_argument("--input", required=True, help="The Excel file produced by 'process'")
    p_verify.add_argument("--report", required=True)
    p_verify.set_defaults(func=_cmd_verify)

    p_finalize = sub.add_parser(
        "finalize", help="Read a labeled Excel file, write the final text,label CSV and its manifest"
    )
    p_finalize.add_argument("--input", required=True)
    p_finalize.add_argument("--output", required=True)
    p_finalize.add_argument("--manifest", required=True)
    p_finalize.add_argument("--version-state", default="data/.corpus_version.json")
    p_finalize.add_argument("--stats", help="stats.json produced by 'process', for the manifest counts")
    p_finalize.set_defaults(func=_cmd_finalize)

    p_pipeline = sub.add_parser(
        "pipeline",
        help="Run everything that doesn't need manual labeling: spec-export + process + verify",
    )
    p_pipeline.add_argument("--spec-output", required=True)
    _add_process_args(p_pipeline)
    p_pipeline.add_argument("--verify-report", required=True)

    return parser


def _cmd_pipeline(args: argparse.Namespace) -> None:
    write_spec(args.spec_output)
    print(f"Wrote anonymization rule spec to {args.spec_output}")
    _cmd_process(args)
    verify_args = argparse.Namespace(input=args.output, report=args.verify_report)
    _cmd_verify(verify_args)


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    if args.command == "pipeline":
        _cmd_pipeline(args)
    else:
        args.func(args)


if __name__ == "__main__":
    main()
